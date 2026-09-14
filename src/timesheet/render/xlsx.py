"""Render a set of days to an .xlsx in the plain 6-column timesheet layout.

Column headings, weekday and month names follow the caller's locale, so the file
lands in whoever has to read it the same language as the live page. All values
are strings: the sheet carries no formulas, and durations and totals are
pre-computed."""
from __future__ import annotations

import io

from .. import i18n
from ..model import Day
from ..timeutil import hm

WIDTHS = [22, 8, 8, 8, 26, 48]


def headers(loc: i18n.Locale) -> list[str]:
    return [loc.col_date, loc.col_from, loc.col_to, loc.col_duration,
            loc.col_project, loc.col_task]


def build_week(week: list[Day], title: str | None = None,
               locale: str | i18n.Locale | None = None) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    loc = locale if isinstance(locale, i18n.Locale) else i18n.get(locale)
    wb = openpyxl.Workbook()
    ws = wb.active
    # Excel rejects a sheet name over 31 characters or containing []:*?/\ — a
    # locale title is safe, but a caller-supplied one need not be.
    ws.title = "".join(c for c in (title or loc.title) if c not in "[]:*?/\\")[:31] or "Sheet1"
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center")
    hdr_font, hdr_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="374151")
    day_font, day_fill = Font(bold=True, color="1F2A44"), PatternFill("solid", fgColor="E8EEF7")

    for c, h in enumerate(headers(loc), 1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill, cell.border = hdr_font, hdr_fill, border

    r, grand = 2, 0
    for day in week:
        grand += day.minutes
        d = day.date
        ws.cell(r, 1, f"{loc.days[d.weekday()]} {d.day} {loc.months[d.month]}"
                      f"  ·  {hm(day.minutes)}{loc.subtitle_hours}")
        for c in range(1, 7):
            ws.cell(r, c).fill, ws.cell(r, c).border = day_fill, border
        ws.cell(r, 1).font = day_font
        r += 1
        for b in day.blocks:
            row = [f"{d.day:02d}-{d.month:02d}-{d.year}", f"{b.start:%H:%M}", f"{b.end:%H:%M}",
                   hm(b.minutes), b.project, b.taak]
            for c, v in enumerate(row, 1):
                cell = ws.cell(r, c, v)
                cell.border = border
                if c in (2, 3, 4):
                    cell.alignment = center
            r += 1

    ws.cell(r + 1, 1, loc.total).font = Font(bold=True)
    tc = ws.cell(r + 1, 4, hm(grand))
    tc.font, tc.alignment = Font(bold=True), center
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
